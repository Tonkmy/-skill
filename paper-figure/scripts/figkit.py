"""figkit - draw editable paper figures as native PowerPoint shapes in the house style.

    import sys; sys.path.insert(0, "<skill>/scripts")
    from figkit import Deck

    deck = Deck()                                   # slide width = full-width canvas (tokens.py)
    fig = deck.figure("column", height=2.4, name="F1 teaser")
    a = fig.box(0.2, 0.5, 1.4, 0.6, "Encoder")                       # grey = existing / frozen
    b = fig.box(2.2, 0.5, 1.4, 0.6, "Adapter", role="orange", emphasis=True)  # the novelty
    fig.arrow(a, b)
    fig.text(2.2, 1.25, 1.4, 0.3, "$e_1 = g(x)$", align="center")
    deck.save("figs.pptx")

Coordinates are inches from the top-left of the figure's export frame and snap to tokens.GRID.
Every shape carries explicit fill / line / font (no theme styles, no shadows), so what LibreOffice
renders, what PowerPoint shows and what the user edits are the same thing.

Text accepts light math markup: anything inside $...$ is set in Inter Italic for letters, with
x_i / x_{ij} subscripts and r^p / r^{p-1} superscripts. Use Unicode for symbols (α ‖ − · ⊕ ≈).
"""
import math
import os
import re

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import tokens as T

EMU_IN = 914400
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_DASH = {"dash": MSO_LINE_DASH_STYLE.DASH, "sysdash": MSO_LINE_DASH_STYLE.SQUARE_DOT, "dot": MSO_LINE_DASH_STYLE.ROUND_DOT}
_SIDE = {"top": 0, "left": 1, "bottom": 2, "right": 3}


def snap(v):
    return round(round(v / T.GRID) * T.GRID, 4)


def _emu(inches):
    return Emu(int(round(inches * EMU_IN)))


def _rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#").upper())


def _color(c):
    """Accept a hex string, a GRAY token name or 'accent.shade' (e.g. 'orange.text')."""
    if c is None:
        return None
    if c in T.GRAY:
        return T.GRAY[c]
    if "." in c:
        name, shade = c.split(".")
        return T.role(name)[shade]
    return c


# ---- rich text --------------------------------------------------------------------------------
_GREEK = "αβγδεζηθικλμνξπρστυφχψωΓΔΘΛΞΠΣΦΨΩ"


def parse_rich(s):
    """Split text into runs: [(text, {'italic': bool, 'baseline': 0|-25|30})].

    Outside $...$ text is upright. Inside, letters are italic; _x / _{..} is a subscript and
    ^x / ^{..} a superscript. In scripts single letters are italic (r^{p−1}: p italic, −1 upright)
    and words stay upright like \\text{} (L_{floor}).
    """
    runs = []
    for i, part in enumerate(re.split(r"(\$[^$]*\$)", s)):
        if not part:
            continue
        if not (part.startswith("$") and part.endswith("$") and len(part) >= 2):
            runs.append((part, {"italic": False, "baseline": 0}))
            continue
        m = part[1:-1]
        j = 0
        while j < len(m):
            ch = m[j]
            if ch in "_^" and j + 1 < len(m):
                base = -25 if ch == "_" else 30
                if m[j + 1] == "{":
                    k = m.index("}", j + 2)
                    script, j = m[j + 2:k], k + 1
                else:
                    script, j = m[j + 1], j + 2
                for tok in re.findall(r"[A-Za-z" + _GREEK + r"]+|[^A-Za-z" + _GREEK + r"]+", script):
                    is_word = tok[0].isalpha() or tok[0] in _GREEK
                    runs.append((tok, {"italic": is_word and len(tok) == 1, "baseline": base}))
                continue
            k = j
            is_letter = ch.isalpha() or ch in _GREEK
            while k < len(m) and m[k] not in "_^" and ((m[k].isalpha() or m[k] in _GREEK) == is_letter):
                k += 1
            runs.append((m[j:k], {"italic": is_letter, "baseline": 0}))
            j = k
    merged = []
    for text, st in runs:
        if merged and merged[-1][1] == st:
            merged[-1] = (merged[-1][0] + text, st)
        else:
            merged.append((text, st))
    return merged


_FONT_CACHE = {}


def text_width(s, size="module", weight="regular"):
    """Rendered width of a (rich) string in canvas inches, measured with the Inter TTFs."""
    pt = T.SIZE.get(size, size)
    width = 0.0
    for line in s.split("\n"):
        w = 0.0
        for text, st in parse_rich(line):
            scale = 0.65 if st["baseline"] else 1.0
            key = (weight, st["italic"])
            if key not in _FONT_CACHE:
                _FONT_CACHE[key] = ImageFont.truetype(T.font_file(weight, st["italic"]), 200)
            w += _FONT_CACHE[key].getlength(text) / 200 * pt * scale
        width = max(width, w)
    return width / 72


def _style_run(run, text, st, size, weight, color, bold=None):
    run.text = text
    f = run.font
    family, is_bold = T.WEIGHTS[weight]
    f.name = family
    f.bold = is_bold if bold is None else bold
    f.italic = st["italic"]
    f.size = Pt(T.SIZE.get(size, size))
    f.color.rgb = _rgb(_color(color))
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):  # pin every script slot to Inter (no theme fallback)
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", family)
    if st["baseline"]:
        rpr.set("baseline", str(st["baseline"] * 1000))


def fill_text(tf, content, size="module", weight="regular", color="text", align="center",
              anchor="middle", pad=None, wrap=False, line_spacing=None):
    """Write (rich, multi-line) content into a text frame with explicit styling."""
    pad = T.PAD if pad is None else pad
    tf.clear()
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = _ANCHOR[anchor]
    tf.margin_left = tf.margin_right = _emu(pad)
    tf.margin_top = tf.margin_bottom = _emu(min(pad, 0.04))
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        if line_spacing:
            p.line_spacing = line_spacing
        for text, st in parse_rich(line) or [("", {"italic": False, "baseline": 0})]:
            _style_run(p.add_run(), text, st, size, weight, color)


# ---- low-level shape styling ------------------------------------------------------------------
def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _set_line(shape, color=None, width=None, dash=None):
    if color is None:
        shape.line.fill.background()
        return
    shape.line.width = Pt(width)
    shape.line.color.rgb = _rgb(_color(color))
    if dash:
        shape.line.dash_style = _DASH[dash]


def _set_fill(shape, color=None):
    if color is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(_color(color))


def _effects(shape, glow=None):
    """Replace effects with nothing (no shadow) or a single soft glow (hex, radius pt, alpha)."""
    sppr = shape._element.spPr
    old = sppr.find(qn("a:effectLst"))
    if old is not None:
        sppr.remove(old)
    eff = etree.SubElement(sppr, qn("a:effectLst"))
    if glow:
        hex_, radius, alpha = glow
        g = etree.SubElement(eff, qn("a:glow"), rad=str(int(Pt(radius))))
        c = etree.SubElement(g, qn("a:srgbClr"), val=_color(hex_))
        etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 100000)))
    # keep schema order: effectLst must precede scene3d / sp3d / extLst
    for tag in ("a:scene3d", "a:sp3d", "a:extLst"):
        el = sppr.find(qn(tag))
        if el is not None:
            sppr.remove(el)
            sppr.append(el)


def _arrowheads(line_shape, head="end"):
    ln = line_shape.line._get_or_add_ln()
    for tag in ("a:headEnd", "a:tailEnd"):
        el = ln.find(qn(tag))
        if el is not None:
            ln.remove(el)
    spec = dict(T.ARROW_HEAD)
    if head in ("start", "both"):
        etree.SubElement(ln, qn("a:headEnd"), **spec)
    if head in ("end", "both"):
        etree.SubElement(ln, qn("a:tailEnd"), **spec)


def _site(node, side):
    """Connection-site index for `side` on the node's shape, or None when it cannot be glued.
    rect / roundRect number their sites top, left, bottom, right (0-3); ellipse has eight,
    counter-clockwise from the top (top 0, left 2, bottom 4, right 6). Groups have none."""
    el = node.shape._element
    if el.tag == qn("p:grpSp"):
        return None
    geom = el.find(".//" + qn("a:prstGeom"))
    prst = geom.get("prst") if geom is not None else None
    if prst == "ellipse":
        return {"top": 0, "left": 2, "bottom": 4, "right": 6}[side]
    if prst in ("rect", "roundRect"):
        return _SIDE[side]
    return None


def _glue(conn, src, src_side, dst, dst_side):
    """Write stCxn / endCxn without moving the ends (they are already at the anchor points)."""
    if isinstance(src, Node) and _site(src, src_side) is not None:
        conn._connect_begin_to(src.shape, _site(src, src_side))
    if isinstance(dst, Node) and _site(dst, dst_side) is not None:
        conn._connect_end_to(dst.shape, _site(dst, dst_side))


class Node:
    """A drawn module: the pptx shape plus anchor points in figure inches."""

    def __init__(self, shape, x, y, w, h):
        self.shape, self.x, self.y, self.w, self.h = shape, x, y, w, h

    left = property(lambda s: (s.x, s.y + s.h / 2))
    right = property(lambda s: (s.x + s.w, s.y + s.h / 2))
    top = property(lambda s: (s.x + s.w / 2, s.y))
    bottom = property(lambda s: (s.x + s.w / 2, s.y + s.h))
    center = property(lambda s: (s.x + s.w / 2, s.y + s.h / 2))

    def side(self, name):
        return getattr(self, name)


# ---- figure (one slide) -----------------------------------------------------------------------
class Figure:
    def __init__(self, deck, slide, slot, width, height, name):
        self.deck, self.slide, self.slot, self.name = deck, slide, slot, name
        self.width, self.height = width, height
        self.ox, self.oy = 0.0, 0.0  # frame origin on the slide
        self._frame = self._add_frame()

    # -- helpers
    @property
    def shapes(self):
        return self.slide.shapes

    def _xy(self, x, y):
        return _emu(self.ox + x), _emu(self.oy + y)

    def _add_frame(self):
        """Invisible rectangle named FRAME: the export bounds that render.py crops to."""
        shp = self.shapes.add_shape(MSO_SHAPE.RECTANGLE, *self._xy(0, 0), _emu(self.width), _emu(self.height))
        _strip_style(shp)
        _set_fill(shp, None)
        _set_line(shp, None)
        _effects(shp)
        shp.name = f"FRAME {self.slot} {self.width:.2f}x{self.height:.2f}in - {self.name}"
        return shp

    def to_back(self, node_or_shape):
        el = getattr(node_or_shape, "shape", node_or_shape)._element
        tree = self.shapes._spTree
        tree.remove(el)
        tree.insert(3, el)  # after nvGrpSpPr, grpSpPr and the FRAME
        return node_or_shape

    # -- modules
    def box(self, x, y, w, h, text=None, role="base", emphasis=False, size="module", weight="medium",
            text_color=None, fill=None, line=None, dashed=False, align="center", anchor="middle",
            pad=None, name=None, radius=None, wrap=False, sub=None):
        """Rounded module. role='base' is grey (existing / frozen); an accent role marks novelty;
        emphasis=True adds the thicker accent outline and a soft glow (use on 1-2 modules).
        sub: optional second line in the note size (a short descriptor, e.g. 'stages 3–4 · head')."""
        x, y, w, h = snap(x), snap(y), snap(w), snap(h)
        c = T.role(role)
        shp = self.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *self._xy(x, y), _emu(w), _emu(h))
        _strip_style(shp)
        r = T.RADIUS if radius is None else radius
        shp.adjustments[0] = min(0.5, r / min(w, h))
        _set_fill(shp, fill if fill is not None else c["fill"])
        lw = T.LINE["emphasis"] if emphasis else T.LINE["border"]
        _set_line(shp, line if line is not None else c["line"], lw, T.DASH if dashed else None)
        _effects(shp, (c["glow"], T.GLOW["radius"], T.GLOW["alpha"]) if emphasis else None)
        if text is not None:
            col = text_color or ("text" if role in (None, "base") else f"{role}.text")
            fill_text(shp.text_frame, text, size, weight, col, align, anchor, pad, wrap)
            if sub:
                p = shp.text_frame.add_paragraph()
                p.alignment = _ALIGN[align]
                sub_col = "muted" if role in (None, "base") else f"{role}.text"
                for t, st in parse_rich(sub):
                    _style_run(p.add_run(), t, st, "note", "regular", sub_col)
        shp.name = name or f"box: {(text or '').splitlines()[0] if text else role}"
        return Node(shp, x, y, w, h)

    def fit(self, text, size="module", weight="medium", pad=None, min_w=0.0):
        """Width (snapped, inches) a box needs to hold `text` on one line."""
        pad = T.PAD if pad is None else pad
        return max(min_w, snap(text_width(text, size, weight) + 2 * pad + 0.06))

    def group(self, x, y, w, h, title=None, role="base", dashed=True, fill="bg", title_size="module",
              title_weight="semibold", name=None):
        """Container for a semantic group, sent behind everything. Title sits inside, top-left."""
        x, y, w, h = snap(x), snap(y), snap(w), snap(h)
        c = T.role(role)
        shp = self.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *self._xy(x, y), _emu(w), _emu(h))
        _strip_style(shp)
        shp.adjustments[0] = min(0.5, T.RADIUS / min(w, h))
        _set_fill(shp, fill if role in (None, "base") else c["fill"] if fill == "bg" else fill)
        _set_line(shp, "guide" if role in (None, "base") else c["line"], T.LINE["border"], T.DASH if dashed else None)
        _effects(shp)
        shp.name = name or f"group: {title or ''}"
        self.to_back(shp)
        node = Node(shp, x, y, w, h)
        if title:
            col = "strong" if role in (None, "base") else f"{role}.text"
            self.text(x + T.GROUP_PAD - T.PAD, y + 0.06, w - 2 * T.GROUP_PAD, 0.4, title, size=title_size,
                      weight=title_weight, color=col, align="left", anchor="top")
        return node

    def text(self, x, y, w, h, content, size="note", weight="regular", color="text", align="left",
             anchor="middle", wrap=False, name=None, pad=None):
        """Free label (no fill, no line). Keep it a label, not a sentence (rule 1)."""
        x, y, w, h = snap(x), snap(y), snap(w), snap(h)
        tb = self.shapes.add_textbox(*self._xy(x, y), _emu(w), _emu(h))
        fill_text(tb.text_frame, content, size, weight, color, align, anchor, T.PAD if pad is None else pad, wrap)
        tb.name = name or f"text: {content.splitlines()[0][:40]}"
        return Node(tb, x, y, w, h)

    def title(self, x, y, label, title=None, w=None):
        """Panel title: '(a)' plus a short title, both in the title size."""
        content = f"{label}  {title}" if title else label
        w = w or text_width(content, "title", "semibold") + 2 * T.PAD + 0.1
        return self.text(x, y, w, 0.45, content, size="title", weight="semibold", color="strong", anchor="top")

    # -- connectors
    def _endpoint(self, ref, side):
        if isinstance(ref, Node):
            return ref.side(side)
        return ref

    def arrow(self, src, dst, src_side="right", dst_side="left", color="text", dashed=False, head="end",
              width=None, name=None):
        """Straight connector. Ends on boxes / operators are glued, so the arrow follows when the user
        moves them; ends on vectors (groups) or points stay where they are drawn."""
        (x1, y1), (x2, y2) = self._endpoint(src, src_side), self._endpoint(dst, dst_side)
        c = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *self._xy(x1, y1), *self._xy(x2, y2))
        _strip_style(c)
        _glue(c, src, src_side, dst, dst_side)
        _set_line(c, color, width or T.LINE["arrow"], T.DASH if dashed else None)
        _arrowheads(c, head)
        c.name = name or "arrow"
        return c

    def elbow(self, src, dst, src_side="right", dst_side="left", color="text", dashed=False, head="end",
              name=None):
        """Horizontal-vertical-horizontal glued connector (for left-to-right flows that change row)."""
        (x1, y1), (x2, y2) = self._endpoint(src, src_side), self._endpoint(dst, dst_side)
        c = self.shapes.add_connector(MSO_CONNECTOR.ELBOW, *self._xy(x1, y1), *self._xy(x2, y2))
        _strip_style(c)
        _glue(c, src, src_side, dst, dst_side)
        _set_line(c, color, T.LINE["arrow"], T.DASH if dashed else None)
        _arrowheads(c, head)
        c.name = name or "elbow"
        return c

    def path(self, points, color="text", dashed=False, head="end", width=None, name=None):
        """Open polyline through `points` (figure inches) for routes connectors cannot express.
        Not glued: if a box moves, its polylines must be redrawn (or edited by hand)."""
        pts = [self._xy(snap(px), snap(py)) for px, py in points]
        fb = self.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
        fb.add_line_segments(pts[1:], close=False)
        shp = fb.convert_to_shape()
        _strip_style(shp)
        _set_fill(shp, None)
        _set_line(shp, color, width or T.LINE["arrow"], T.DASH if dashed else None)
        _effects(shp)
        _arrowheads(shp, head)
        shp.name = name or "path"
        return shp

    def dot(self, x, y, d=0.08, color="text"):
        """Junction dot where a flow splits."""
        shp = self.shapes.add_shape(MSO_SHAPE.OVAL, *self._xy(x - d / 2, y - d / 2), _emu(d), _emu(d))
        _strip_style(shp)
        _set_fill(shp, color)
        _set_line(shp, None)
        _effects(shp)
        shp.name = "dot"
        return Node(shp, x - d / 2, y - d / 2, d, d)

    def op(self, x, y, symbol="+", d=0.36, role="base", emphasis=False, name=None):
        """Operator node (⊕, ×, ·) centred at (x, y)."""
        c = T.role(role)
        shp = self.shapes.add_shape(MSO_SHAPE.OVAL, *self._xy(x - d / 2, y - d / 2), _emu(d), _emu(d))
        _strip_style(shp)
        _set_fill(shp, "white" if role in (None, "base") else c["fill"])
        _set_line(shp, "text" if role in (None, "base") else c["line"], T.LINE["emphasis"] if emphasis else T.LINE["border"])
        _effects(shp, (c["glow"], T.GLOW["radius"], T.GLOW["alpha"]) if emphasis else None)
        fill_text(shp.text_frame, symbol, "title", "medium", "text" if role in (None, "base") else f"{role}.text",
                  "center", "middle", pad=0.0)
        shp.name = name or f"op {symbol}"
        return Node(shp, x - d / 2, y - d / 2, d, d)

    def vector(self, x, y, n=4, cell=0.16, role="base", emphasis=False, horizontal=False, name=None):
        """Embedding icon: a column (or row) of n small rounded cells, grouped."""
        c = T.role(role)
        grp = self.shapes.add_group_shape()
        gap = cell * 0.25
        for i in range(n):
            cx = x + (i * (cell + gap) if horizontal else 0)
            cy = y + (0 if horizontal else i * (cell + gap))
            s = grp.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *self._xy(cx, cy), _emu(cell), _emu(cell))
            _strip_style(s)
            s.adjustments[0] = 0.25
            _set_fill(s, c["fill"])
            _set_line(s, c["line"], T.LINE["emphasis"] if emphasis else T.LINE["border"])
            _effects(s)
        grp.name = name or "vector"
        w = n * cell + (n - 1) * gap if horizontal else cell
        h = cell if horizontal else n * cell + (n - 1) * gap
        return Node(grp, x, y, w, h)

    # -- icons and pictures
    def icon(self, kind, x, y, size=0.26, color=None):
        """'snowflake' (frozen, grey by default) or 'flame' (trainable, orange by default)."""
        cx, cy, r = x + size / 2, y + size / 2, size / 2
        if kind == "snowflake":
            col = color or "muted"
            grp = self.shapes.add_group_shape()
            segs = []
            for k in range(3):
                a = math.pi / 3 * k
                segs.append(((cx - r * math.cos(a), cy - r * math.sin(a)), (cx + r * math.cos(a), cy + r * math.sin(a))))
            for k in range(6):
                a = math.pi / 3 * k
                px, py = cx + 0.62 * r * math.cos(a), cy + 0.62 * r * math.sin(a)
                for s in (-1, 1):
                    b = a + s * math.pi / 4
                    segs.append(((px, py), (px + 0.32 * r * math.cos(b), py + 0.32 * r * math.sin(b))))
            for (x1, y1), (x2, y2) in segs:
                ln = grp.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *self._xy(x1, y1), *self._xy(x2, y2))
                _strip_style(ln)
                _set_line(ln, col, max(1.0, size * 5.5))
                ln.line._get_or_add_ln().set("cap", "rnd")
            grp.name = "icon: snowflake (frozen)"
            return Node(grp, x, y, size, size)
        if kind == "flame":
            col = color or "orange.line"
            grp = self.shapes.add_group_shape()
            for scale, dy, c in ((1.0, 0.0, col), (0.5, 0.42, "orange.glow")):
                rr = r * scale
                ccx, ccy = cx, cy + dy * r
                pts = [self._xy(ccx + 0.8 * rr * math.sin(t) * math.sin(t / 2), ccy - rr * math.cos(t))
                       for t in (2 * math.pi * i / 48 for i in range(49))]
                fb = grp.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
                fb.add_line_segments(pts[1:], close=True)
                shp = fb.convert_to_shape()
                _strip_style(shp)
                _set_fill(shp, c)
                _set_line(shp, None)
                _effects(shp)
            grp.name = "icon: flame (trainable)"
            return Node(grp, x, y, size, size)
        raise ValueError(kind)

    def picture(self, path, x, y, w=None, h=None, rounded=False, border=None, name=None):
        """Insert an image (plot PNG, face crops). rounded=True clips to the house corner radius."""
        pic = self.shapes.add_picture(path, *self._xy(snap(x), snap(y)),
                                      _emu(w) if w else None, _emu(h) if h else None)
        if rounded:
            geom = pic._element.spPr.find(qn("a:prstGeom"))
            geom.set("prst", "roundRect")
            av = geom.find(qn("a:avLst"))
            if av is None:
                av = etree.SubElement(geom, qn("a:avLst"))
            wi, hi = pic.width / EMU_IN, pic.height / EMU_IN
            etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(min(0.5, T.RADIUS / min(wi, hi)) * 100000)}")
        if border:
            _set_line(pic, border, T.LINE["border"])
        pic.name = name or f"picture: {os.path.basename(path)}"
        return Node(pic, snap(x), snap(y), pic.width / EMU_IN, pic.height / EMU_IN)

    # -- legend
    def legend(self, x, y, items, gap=0.45, size="note"):
        """One row of legend entries. items: (kind, label) with kind in
        'arrow' | 'dashed' | 'role:<name>' | 'emphasis:<name>' | 'icon:<kind>'."""
        cx = x
        h = 0.3
        for kind, label in items:
            if kind in ("arrow", "dashed"):
                self.arrow((cx, y + h / 2), (cx + 0.45, y + h / 2), dashed=kind == "dashed")
                cx += 0.55
            elif kind.startswith(("role:", "emphasis:")):
                k, r = kind.split(":")
                self.box(cx, y + 0.03, 0.35, h - 0.06, role=r, emphasis=k == "emphasis", name=f"legend swatch {r}")
                cx += 0.45
            elif kind.startswith("icon:"):
                self.icon(kind.split(":")[1], cx, y + 0.02, size=h - 0.04)
                cx += 0.35
            tw = text_width(label, size, "regular") + 2 * T.PAD + 0.05
            self.text(cx - T.PAD + 0.02, y, tw, h, label, size=size, color="muted")
            cx += tw - T.PAD + gap


# ---- deck -------------------------------------------------------------------------------------
class Deck:
    """One figure per slide. A new deck's slide width is the full-width canvas; an existing deck
    (path=...) keeps its slide size, which must be at least that wide."""

    def __init__(self, path=None, venue="ieee", height=None):
        self.venue = venue
        self.prs = Presentation(path) if path else Presentation()
        full = T.frame_width("full", venue)
        if not path:
            self.prs.slide_width = _emu(full)
            self.prs.slide_height = _emu(height or 6.0)
        self._auto_height = path is None and height is None
        self._max_h = 0.0
        layouts = self.prs.slide_layouts
        self._blank = min(layouts, key=lambda l: len(l.placeholders))

    def figure(self, slot="full", height=4.0, name="figure", width=None, origin=None, notes=None):
        """Add one slide holding one figure. origin='center' centres the export frame on the
        slide (useful in an existing deck with large slides); notes go to the speaker notes."""
        width = width or T.frame_width(slot, self.venue)
        slide = self.prs.slides.add_slide(self._blank)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        self._max_h = max(self._max_h, height)
        fig = Figure.__new__(Figure)
        fig.deck, fig.slide, fig.slot, fig.name, fig.width, fig.height = self, slide, slot, name, width, height
        fig.ox, fig.oy = 0.0, 0.0
        if origin == "center":
            fig.ox = snap((self.prs.slide_width / EMU_IN - width) / 2)
            fig.oy = snap((self.prs.slide_height / EMU_IN - height) / 2)
        elif origin:
            fig.ox, fig.oy = origin
        fig._frame = fig._add_frame()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return fig

    def save(self, path, force=False):
        """Write the deck. Refuses when Office has the target open (a ~$ lock file sits next to it):
        a later save in PowerPoint would silently drop these slides."""
        lock = os.path.join(os.path.dirname(os.path.abspath(path)), "~$" + os.path.basename(path))
        if os.path.exists(lock) and not force:
            raise RuntimeError(f"{os.path.basename(path)} is open in PowerPoint ({lock} exists): "
                               "ask the user to save and close it first")
        if self._auto_height:
            self.prs.slide_height = _emu(snap(self._max_h))
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        return path
