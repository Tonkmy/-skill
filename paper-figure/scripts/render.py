"""Render a deck to PNG for self-checking, and crop figures to their FRAME for LaTeX.

    python render.py deck.pptx -o out/                    # PNG of every slide
    python render.py deck.pptx -o out/ --frame            # crop to FRAME -> slideN.png + slideN.pdf
    python render.py exported.pdf --pptx deck.pptx -o out/ --frame   # crop a PowerPoint "Save as PDF"
    options: --slides 1,3   --dpi 220

LibreOffice renders the preview. In headless mode it only sees fonts listed by fontconfig, so this
script points it at Homebrew's fonts.conf (which covers ~/Library/Fonts); without that, Inter,
Arial and Cambria silently fall back to bundled fonts. PowerPoint stays the reference for the
final look (glow and kerning can differ slightly).
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile

import pymupdf
from pptx import Presentation

FONTCONF = ["/opt/homebrew/etc/fonts/fonts.conf", "/usr/local/etc/fonts/fonts.conf"]


def soffice_pdf(pptx, outdir):
    exe = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    env = dict(os.environ)
    for conf in FONTCONF:
        if os.path.exists(conf):
            env.setdefault("FONTCONFIG_FILE", conf)
            break
    profile = os.path.join(tempfile.gettempdir(), "paper-figure-lo-profile")
    cmd = [exe, "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", outdir, pptx]
    subprocess.run(cmd, env=env, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(pptx))[0] + ".pdf")
    if not os.path.exists(pdf):
        sys.exit(f"LibreOffice did not produce {pdf}")
    return pdf


def frames(pptx):
    """{slide index: (x, y, w, h) in EMU} of the shape whose name starts with FRAME, plus slide width."""
    prs = Presentation(pptx)
    out = {}
    for i, slide in enumerate(prs.slides):
        for shp in slide.shapes:
            if shp.name.startswith("FRAME"):
                out[i] = (shp.left, shp.top, shp.width, shp.height)
                break
    return out, prs.slide_width


def crop_pdf(doc, i, clip, path):
    """Page i of doc as a stand-alone figure PDF for LaTeX. Content outside the frame (the working
    note, anything parked off-canvas) is deleted rather than hidden, since pdflatex embeds the whole
    page; document metadata is cleared; MediaBox = CropBox = the frame. Text and lines stay vector."""
    one = pymupdf.open()
    one.insert_pdf(doc, from_page=i, to_page=i)
    page, r, e = one[0], one[0].rect, 0.5  # e: keep the cut 0.5 pt clear of the frame edge
    for box in (pymupdf.Rect(r.x0, r.y0, r.x1, clip.y0 - e), pymupdf.Rect(r.x0, clip.y1 + e, r.x1, r.y1),
                pymupdf.Rect(r.x0, r.y0, clip.x0 - e, r.y1), pymupdf.Rect(clip.x1 + e, r.y0, r.x1, r.y1)):
        if not box.is_empty:
            page.add_redact_annot(box, fill=False)
    page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS,
                          graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED)
    page.clean_contents(sanitize=True)  # also drops resources nothing uses now (e.g. the note's CJK font)
    page.set_cropbox(clip)
    one.xref_set_key(page.xref, "MediaBox", one.xref_get_key(page.xref, "CropBox")[1])
    one.set_metadata({})
    one.del_xml_metadata()
    one.save(path, garbage=4, deflate=True)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src", help=".pptx (rendered with LibreOffice) or .pdf (e.g. PowerPoint export)")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--pptx", help="deck that holds the FRAME shapes when src is a PDF")
    ap.add_argument("--frame", action="store_true", help="crop each slide to its FRAME and write a PDF too")
    ap.add_argument("--slides", help="1-based list, e.g. 1,3")
    ap.add_argument("--dpi", type=int, default=200)
    a = ap.parse_args()

    os.makedirs(a.out, exist_ok=True)
    src = os.path.abspath(a.src)
    deck = src if src.endswith(".pptx") else (os.path.abspath(a.pptx) if a.pptx else None)
    pdf = soffice_pdf(src, a.out) if src.endswith(".pptx") else src
    doc = pymupdf.open(pdf)
    fr, slide_w = frames(deck) if (a.frame and deck) else ({}, None)
    if a.frame and not deck:
        sys.exit("--frame on a PDF needs --pptx to know where the FRAMEs are")
    pick = [int(s) - 1 for s in a.slides.split(",")] if a.slides else range(doc.page_count)
    stem = os.path.splitext(os.path.basename(src))[0]
    for i in pick:
        page = doc[i]
        clip = None
        if a.frame and i in fr:
            k = page.rect.width / slide_w
            x, y, w, h = (v * k for v in fr[i])
            clip = pymupdf.Rect(x, y, x + w, y + h) & page.rect
            out = crop_pdf(doc, i, clip, os.path.join(a.out, f"{stem}_slide{i + 1}.pdf"))
            page = pymupdf.open(out)[0]  # preview what LaTeX will get
            clip = None
        png = os.path.join(a.out, f"{stem}_slide{i + 1}.png")
        page.get_pixmap(dpi=a.dpi, clip=clip).save(png)
        print(png)


if __name__ == "__main__":
    main()
