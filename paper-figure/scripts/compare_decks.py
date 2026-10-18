"""Check that every slide of the old deck survives unchanged in the new one (C14N), in order.

    python compare_decks.py before.pptx after.pptx

Slides are matched in order by content, so slides inserted anywhere (and python-pptx renumbering
slideN.xml on save) do not matter. Each slide is compared by its XML, its notes, its images and other
related parts; masters, layouts and themes are compared too. Exit code 1 when an old slide is missing
or changed.
"""
import argparse
import hashlib
import sys

from lxml import etree
from pptx import Presentation


def c14n(el):
    return etree.tostring(el, method="c14n")


def slide_sig(slide):
    h = hashlib.sha256(c14n(slide._element))
    for rid, rel in sorted(slide.part.rels.items()):
        kind = rel.reltype.rsplit("/", 1)[-1]
        if rel.is_external:
            h.update(f"{rid}:{kind}:{rel.target_ref}".encode())
            continue
        part = rel.target_part
        if kind == "slideLayout":
            body = str(part.partname).encode()  # shared; compared below with the masters
        elif kind == "notesSlide":
            body = c14n(part._element)
        else:
            body = part.blob
        h.update(f"{rid}:{kind}:".encode() + hashlib.sha256(body).digest())
    return h.hexdigest()


def shared_parts(prs):
    out = {}
    for part in prs.part.package.iter_parts():
        name = str(part.partname)
        if any(k in name for k in ("slideMaster", "slideLayout", "theme")):
            out[name] = hashlib.sha256(c14n(part._element) if hasattr(part, "_element") else part.blob).hexdigest()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("before")
    ap.add_argument("after")
    a = ap.parse_args()
    old, new = Presentation(a.before), Presentation(a.after)
    so, sn = [slide_sig(s) for s in old.slides], [slide_sig(s) for s in new.slides]
    matched, missing, j = {}, [], 0
    for i, sig in enumerate(so, 1):
        k = next((k for k in range(j, len(sn)) if sn[k] == sig), None)
        if k is None:
            missing.append(i)
            continue
        matched[i], j = k + 1, k + 1
    added = [k for k in range(1, len(sn) + 1) if k not in matched.values()]
    for i in range(1, len(so) + 1):
        print(f"old {i:2d} -> " + (f"new {matched[i]:2d}  unchanged" if i in matched else "CHANGED or missing"))
    print(f"added in new: {added or 'none'}")
    same_shared = shared_parts(old) == shared_parts(new)
    print(f"masters / layouts / themes: {'unchanged' if same_shared else 'CHANGED'}")
    sys.exit(1 if missing or not same_shared else 0)


if __name__ == "__main__":
    main()
