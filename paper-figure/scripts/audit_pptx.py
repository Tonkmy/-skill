"""Audit a deck against the figure rules (../style-rules.md) and print a Markdown report.

    python audit_pptx.py deck.pptx [--slides 1,2]

Checks what can be measured: sentences in labels (R1), grey vs colour share (R2), typefaces and
sizes (R3), rounded shapes / radii / line and arrow widths (R4), number of hues (R5), labels that
collide or leave the frame, lines drawn through label text and neighbouring modules closer than
tokens.GAP (R6). Judgement calls (is the structure clear?) still need eyes on the rendered PNG.
"""
import argparse
import colorsys
import re
from collections import Counter

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

import tokens as T

EMU_IN = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
SERIF = ("times", "cambria", "georgia", "garamond", "palatino", "book antiqua", "baskerville", "charter",
         "libertine", "computer modern", "cmu serif", "latin modern", "stix", "caladea", "songti", "simsun",
         "宋体", "minion", "constantia", "linux libertine", "noto serif", "source serif", "dejavu serif")
SANS = ("inter", "arial", "helvetica", "calibri", "segoe", "roboto", "open sans", "lato", "source sans",
        "sf pro", "pingfang", "yahei", "微软雅黑", "noto sans", "dejavu sans", "avenir", "futura", "gill sans",
        "verdana", "tahoma", "liberation sans", "myriad", "frutiger", "ibm plex sans", "manrope")


def classify(face):
    f = (face or "").lower()
    if any(s in f for s in SERIF):
        return "serif"
    if any(s in f for s in SANS):
        return "sans"
    if "math" in f or "symbol" in f:
        return "symbol"
    return "?"


def theme_fonts(prs):
    """Theme fonts (+mj-lt / +mn-lt) and theme colours (scheme:accent1 -> hex) of the first master."""
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        out = {"+mj-lt": root.find(f".//{A}majorFont/{A}latin").get("typeface"),
               "+mn-lt": root.find(f".//{A}minorFont/{A}latin").get("typeface")}
        scheme = root.find(f".//{A}clrScheme")
        for el in scheme if scheme is not None else []:
            name = etree.QName(el).localname
            c = el.find(f"{A}srgbClr")
            sysc = el.find(f"{A}sysClr")
            val = c.get("val") if c is not None else (sysc.get("lastClr") if sysc is not None else None)
            if val:
                out["scheme:" + name] = val.upper()
        for alias, src in (("tx1", "dk1"), ("bg1", "lt1"), ("tx2", "dk2"), ("bg2", "lt2")):
            if "scheme:" + src in out:
                out["scheme:" + alias] = out["scheme:" + src]
        return out
    except Exception:
        return {}


def walk(shapes):
    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(s.shapes)


def _clr(parent):
    """Colour inside a fill-bearing element: hex, 'scheme:x', 'gradient' or None."""
    if parent is None:
        return None
    sf = parent.find(f"{A}solidFill")
    if sf is not None:
        c = sf.find(f"{A}srgbClr")
        if c is not None:
            return c.get("val").upper()
        c = sf.find(f"{A}schemeClr")
        if c is not None:
            return "scheme:" + c.get("val")
        return "other"
    if parent.find(f"{A}gradFill") is not None:
        return "gradient"
    return None


def shape_info(s, theme):
    el = s._element
    sppr = el.find(qn("p:spPr"))
    style = el.find(qn("p:style"))
    info = {"name": s.name, "geom": None, "adj": None, "fill": None, "line": None, "lw": None, "dash": None,
            "heads": (), "glow": False, "runs": [], "text": "", "kind": s.shape_type,
            "box": (s.left or 0, s.top or 0, s.width or 0, s.height or 0)}
    if sppr is not None:
        g = sppr.find(f"{A}prstGeom")
        if g is not None:
            info["geom"] = g.get("prst")
            gd = g.find(f"{A}avLst/{A}gd")
            if gd is not None and gd.get("fmla", "").startswith("val "):
                info["adj"] = int(gd.get("fmla")[4:]) / 100000
        elif sppr.find(f"{A}custGeom") is not None:
            info["geom"] = "custom"
        info["fill"] = _clr(sppr)
        if info["fill"] is None and sppr.find(f"{A}noFill") is None and style is not None:
            ref = style.find(f"{A}fillRef")
            if ref is not None and ref.get("idx", "0") != "0":
                info["fill"] = "theme"
        ln = sppr.find(f"{A}ln")
        if ln is not None and ln.find(f"{A}noFill") is None:
            info["line"] = _clr(ln) or ("theme" if style is not None else None)
            info["lw"] = round(int(ln.get("w", 9525)) / 12700, 2)
            d = ln.find(f"{A}prstDash")
            info["dash"] = d.get("val") if d is not None else None
            info["heads"] = tuple((h.get("type") or "") for h in (ln.find(f"{A}headEnd"), ln.find(f"{A}tailEnd"))
                                  if h is not None and h.get("type") not in (None, "none"))
        elif ln is None and style is not None:
            ref = style.find(f"{A}lnRef")
            if ref is not None and ref.get("idx", "0") != "0":
                info["line"], info["lw"] = "theme", 0.75
        eff = sppr.find(f"{A}effectLst")
        info["glow"] = eff is not None and eff.find(f"{A}glow") is not None
    if getattr(s, "has_text_frame", False) and s.has_text_frame:
        info["text"] = s.text_frame.text.strip()
        for p in s.text_frame.paragraphs:
            ppr = p._p.find(f"{A}pPr")
            dflt = ppr.find(f"{A}defRPr") if ppr is not None else None
            for r in p.runs:
                if not r.text.strip():
                    continue
                rpr = r._r.find(f"{A}rPr")
                face, sz, base = None, None, 0
                for src in (rpr, dflt):
                    if src is None:
                        continue
                    lat = src.find(f"{A}latin")
                    face = face or (lat.get("typeface") if lat is not None else None)
                    sz = sz or (int(src.get("sz")) / 100 if src.get("sz") else None)
                if rpr is not None and rpr.get("baseline"):
                    base = int(rpr.get("baseline"))
                if face is None or face.startswith("+"):
                    face = theme.get(face or "+mn-lt", "theme") + " (theme)"
                col = _clr(rpr) if rpr is not None else None
                info["runs"].append({"face": face, "size": sz or 18.0, "baseline": base, "color": col, "text": r.text})
    return info


_PALETTE = {v.upper(): name for name, shades in T.ACCENTS.items() for v in shades.values()}
_PALETTE.update({v.upper(): None for v in T.GRAY.values()})


def hue_family(hex_):
    """Accent name for house-palette colours; otherwise a hue name, or None for neutrals
    (HLS saturation, so pale tints still count)."""
    if not hex_ or len(hex_) != 6:
        return None
    if hex_.upper() in _PALETTE:
        return _PALETTE[hex_.upper()]
    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if max(r, g, b) - min(r, g, b) < 0.03 or s < 0.25:
        return None  # grey / black / white
    names = ["red", "orange", "yellow", "lime", "green", "teal", "cyan", "azure", "blue", "violet", "purple", "pink"]
    return names[int(((h * 360 + 15) % 360) // 30)]


# ---- geometry for R6: absolute boxes, line segments, text ink, module gaps --------------------
def walk_abs(shapes, tf=(0, 1, 0, 1)):
    """Yield (shape, (x, y, w, h) in slide EMU), resolving group child transforms."""
    a, b, c, d = tf
    for s in shapes:
        x, y, w, h = s.left or 0, s.top or 0, s.width or 0, s.height or 0
        box = (a + x * b, c + y * d, w * b, h * d)
        yield s, box
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = s._element.find(qn("p:grpSpPr")).find(f"{A}xfrm")
            off = xfrm.find(f"{A}chOff") if xfrm is not None else None
            ext = xfrm.find(f"{A}chExt") if xfrm is not None else None
            chx, chy = (int(off.get("x")), int(off.get("y"))) if off is not None else (x, y)
            cw, ch = (int(ext.get("cx")), int(ext.get("cy"))) if ext is not None else (w, h)
            sx, sy = (w / cw if cw else 1), (h / ch if ch else 1)
            yield from walk_abs(s.shapes, (box[0] - chx * sx * b, b * sx, box[1] - chy * sy * d, d * sy))


def segments(s, box):
    """Straight segments of a visible connector or open freeform (figkit path), in slide EMU."""
    el = s._element
    sppr = el.find(qn("p:spPr"))
    if sppr is None:
        return []
    ln = sppr.find(f"{A}ln")
    if ln is not None and ln.find(f"{A}noFill") is not None:
        return []
    xfrm = sppr.find(f"{A}xfrm")
    fh = xfrm is not None and xfrm.get("flipH") == "1"
    fv = xfrm is not None and xfrm.get("flipV") == "1"
    x, y, w, h = box
    if el.tag == qn("p:cxnSp"):
        return [((x + w if fh else x, y + h if fv else y), (x if fh else x + w, y if fv else y + h))]
    cust = sppr.find(f"{A}custGeom")
    if cust is None or sppr.find(f"{A}solidFill") is not None or sppr.find(f"{A}gradFill") is not None:
        return []  # not a freeform, or a filled one (icons)
    out = []
    for path in cust.iter(f"{A}path"):
        pw, ph = int(path.get("w", 0)) or 1, int(path.get("h", 0)) or 1
        prev = None
        for cmd in path:
            pts = cmd.findall(f"{A}pt")
            if not pts:
                continue
            px, py = int(pts[-1].get("x")), int(pts[-1].get("y"))
            p = (x + (pw - px if fh else px) * w / pw, y + (ph - py if fv else py) * h / ph)
            if etree.QName(cmd).localname != "moveTo" and prev is not None:
                out.append((prev, p))  # curves approximated by their chord
            prev = p
    return out


_FONT_CACHE = {}


def _run_width(text, face, size, italic, bold, baseline):
    """Advance width of a run in EMU, measured with the Inter TTFs (other faces approximated by Inter)."""
    f = (face or "").lower()
    weight = "semibold" if "semibold" in f else "medium" if "medium" in f else "bold" if bold else "regular"
    key = (weight, italic)
    if key not in _FONT_CACHE:
        try:
            _FONT_CACHE[key] = ImageFont.truetype(T.font_file(weight, italic), 200)
        except (OSError, FileNotFoundError):
            _FONT_CACHE[key] = None
    font = _FONT_CACHE[key]
    pt = (font.getlength(text) / 200 * size) if font else 0.55 * size * len(text)
    return pt * (0.65 if baseline else 1.0) / 72 * EMU_IN


def ink_rects(s, box):
    """Rough glyph area of each text line (alignment, anchor and insets applied), in slide EMU."""
    if not (getattr(s, "has_text_frame", False) and s.has_text_frame and s.text_frame.text.strip()):
        return []
    body = s.text_frame._txBody.find(f"{A}bodyPr")
    ins = {k: int(body.get(k, dv)) for k, dv in (("lIns", 91440), ("rIns", 91440), ("tIns", 45720), ("bIns", 45720))}
    lines = []
    for p in s.text_frame.paragraphs:
        ppr = p._p.find(f"{A}pPr")
        algn = ppr.get("algn", "l") if ppr is not None else "l"
        width, size = 0.0, 0.0
        for r in p.runs:
            rpr = r._r.find(f"{A}rPr")
            get = (lambda k, dv=None: rpr.get(k, dv)) if rpr is not None else (lambda k, dv=None: dv)
            sz = int(get("sz")) / 100 if get("sz") else 18.0
            lat = rpr.find(f"{A}latin") if rpr is not None else None
            base = int(get("baseline", 0))
            width += _run_width(r.text, lat.get("typeface") if lat is not None else "Inter", sz, get("i") == "1",
                                get("b") == "1", base)
            if not base:
                size = max(size, sz)
        lines.append((algn, width, size or 12.0))
    x, y, w, h = box
    lh = lambda sz: sz * 1.2 / 72 * EMU_IN  # noqa: E731
    block = sum(lh(sz) for _, _, sz in lines)
    anchor = body.get("anchor", "t")
    top = {"t": y + ins["tIns"], "b": y + h - ins["bIns"] - block}.get(
        anchor, y + ins["tIns"] + (h - ins["tIns"] - ins["bIns"] - block) / 2)
    rects = []
    for algn, width, sz in lines:
        if width > 0:
            x0 = {"ctr": x + ins["lIns"] + (w - ins["lIns"] - ins["rIns"] - width) / 2,
                  "r": x + w - ins["rIns"] - width}.get(algn, x + ins["lIns"])
            rects.append((x0, top + 0.2 * lh(sz), width, 0.6 * lh(sz)))
        top += lh(sz)
    return rects


def _hits(seg, rect):
    """Does the segment cross the rectangle? (Liang-Barsky clipping)"""
    (x1, y1), (x2, y2) = seg
    rx, ry, rw, rh = rect
    dx, dy = x2 - x1, y2 - y1
    t0, t1 = 0.0, 1.0
    for p, q in ((-dx, x1 - rx), (dx, rx + rw - x1), (-dy, y1 - ry), (dy, ry + rh - y1)):
        if p == 0:
            if q < 0:
                return False
        else:
            t = q / p
            if p < 0:
                t0 = max(t0, t)
            else:
                t1 = min(t1, t)
            if t0 > t1:
                return False
    return True


def lines_through_text(slide):
    """Labels that a connector or path is drawn through (rule 1: lines never cross text)."""
    items = [(s, box) for s, box in walk_abs(slide.shapes) if not s.name.startswith(("FRAME", "NOTE"))]
    segs = [sg for s, box in items for sg in segments(s, box)]
    hit = []
    for s, box in items:
        if any(_hits(sg, r) for r in ink_rects(s, box) for sg in segs):
            hit.append(s.text_frame.text.strip().replace("\n", " ")[:24])
    return hit


def tight_gaps(slide):
    """Neighbouring modules (same row or column) closer than tokens.GAP. Modules are top-level boxes,
    ellipses, pictures and groups (stacks, vectors) at least 0.3 in long; containers are skipped."""
    mods = []
    for s in slide.shapes:
        if s.name.startswith(("FRAME", "NOTE")) or s.shape_type == MSO_SHAPE_TYPE.LINE:
            continue
        if s._element.tag == qn("p:cxnSp") or max(s.width or 0, s.height or 0) < 0.3 * EMU_IN:
            continue
        if s.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP):
            sppr = s._element.find(qn("p:spPr"))
            if sppr is None or sppr.find(f"{A}custGeom") is not None:
                continue
            ln = sppr.find(f"{A}ln")
            visible = sppr.find(f"{A}solidFill") is not None or (ln is not None and ln.find(f"{A}noFill") is None
                                                                 and ln.find(f"{A}solidFill") is not None)
            if not visible:
                continue  # plain text box
        mods.append((s.name, (s.left, s.top, s.width, s.height)))
    tol = 0.01 * EMU_IN
    inside = lambda p, q: (p[0] >= q[0] - tol and p[1] >= q[1] - tol and p[0] + p[2] <= q[0] + q[2] + tol  # noqa: E731
                           and p[1] + p[3] <= q[1] + q[3] + tol)
    mods = [m for m in mods if not any(o is not m and inside(o[1], m[1]) for o in mods)]  # drop containers
    out, lim = [], (T.GAP - 0.02) * EMU_IN
    for i in range(len(mods)):
        for j in range(i + 1, len(mods)):
            (na, a), (nb, b) = mods[i], mods[j]
            for k in (0, 1):  # k = 0: same row, horizontal gap; k = 1: same column, vertical gap
                o = 1 - k
                overlap = min(a[o] + a[o + 2], b[o] + b[o + 2]) - max(a[o], b[o])
                gap = max(a[k], b[k]) - min(a[k] + a[k + 2], b[k] + b[k + 2])
                if overlap >= 0.5 * min(a[o + 2], b[o + 2]) and 0 < gap < lim:
                    out.append(f"{na[:16]}↔{nb[:16]} {gap / EMU_IN:.2f}")
    return out


def sentence_like(t):
    t = " ".join(t.split())
    words = [w for w in re.findall(r"[A-Za-z一-鿿]+", t) if len(w) > 1]  # skip math variables
    # a sentence boundary is ". " + capital; "Aux. tail" (abbreviation + lower case) is a label
    return len(words) > 6 or bool(re.search(r"[a-z]\.\s+[A-Z]", t)) or (t.endswith(".") and not re.search(r"\d\.$", t))


def audit_slide(slide, idx, theme):
    shapes = [shape_info(s, theme) for s in walk(slide.shapes)]
    for x in shapes:  # theme colours -> hex
        for k in ("fill", "line"):
            if isinstance(x[k], str) and x[k].startswith("scheme:"):
                x[k] = theme.get(x[k], x[k])
        for r in x["runs"]:
            if isinstance(r["color"], str) and r["color"].startswith("scheme:"):
                r["color"] = theme.get(r["color"], r["color"])
    frame = next((x for x in shapes if x["name"].startswith("FRAME")), None)
    shapes = [x for x in shapes if not x["name"].startswith(("FRAME", "NOTE"))]  # notes are not exported
    rows, facts = [], {}
    facts["frame"] = frame["name"] if frame else "（无 FRAME：整页即图）"
    facts["pictures"] = sum(1 for x in shapes if x["kind"] == MSO_SHAPE_TYPE.PICTURE)
    if all(x["kind"] == MSO_SHAPE_TYPE.PICTURE for x in shapes):
        rows.append(("全部规则", "—", "整页只有图片（如 matplotlib 导出），矢量审计不适用：看渲染 PNG，并检查出图脚本是否用 plot_style.py"))
        return rows, facts

    # R1 sentences
    texts = [x for x in shapes if x["text"]]
    bad = [x["text"].replace("\n", " ")[:70] for x in texts if sentence_like(x["text"])]
    rows.append(("R1 结构不是句子", "⚠️" if bad else "✅",
                 f"{len(bad)} 个文本框像句子：" + "；".join(f"「{b}」" for b in bad[:4]) if bad else f"{len(texts)} 个文本框都是短标签"))

    # R2 grey base vs colour
    mods = [x for x in shapes if x["geom"] not in (None, "line", "ellipse", "custom") and x["fill"]
            and len(x["fill"]) == 6 and x["fill"] != "FFFFFF" and x["box"][2] > 0.25 * EMU_IN]
    area = lambda x: x["box"][2] * x["box"][3]
    col_area = sum(area(x) for x in mods if hue_family(x["fill"]))
    tot = sum(area(x) for x in mods) or 1
    share = col_area / tot
    emph = sum(1 for x in shapes if x["glow"] or (x["lw"] or 0) >= T.LINE["emphasis"] - 0.01)
    r2_bad = share > 0.5 or (len(mods) >= 3 and emph == 0)
    rows.append(("R2 创新点高亮", "⚠️" if r2_bad else "✅",
                 f"彩色模块占模块面积 {share:.0%}（基础组件应为灰，宜 ≤50%）；强调描边/glow 元素 {emph} 个" +
                 ("（没有强调的创新模块）" if len(mods) >= 3 and emph == 0 else "")))

    # R3 fonts and sizes
    runs = [r for x in shapes for r in x["runs"]]
    faces = Counter(r["face"] for r in runs)
    classes = {classify(f.replace(" (theme)", "")) for f in faces}
    sizes = sorted({r["size"] for r in runs if not r["baseline"]})
    non_inter = [f for f in faces if not f.lower().startswith("inter")]
    tiers = set(T.SIZE.values())
    st = "❌" if {"serif", "sans"} <= classes or len(sizes) > 3 else ("⚠️" if non_inter or set(sizes) - tiers else "✅")
    rows.append(("R3 字体统一", st, "字体 " + "，".join(f"{f}×{n}" for f, n in faces.most_common()) +
                 f"；字号 {', '.join(f'{s:g}' for s in sizes)} pt" + ("；混用衬线与无衬线" if {"serif", "sans"} <= classes else "")))

    # R4 shapes and lines
    sharp = [x for x in shapes if x["geom"] in ("rect", "cube", "trapezoid") and (x["fill"] or x["line"])]
    radii = sorted({round((x["adj"] if x["adj"] is not None else 0.16667) * min(x["box"][2], x["box"][3]) / EMU_IN, 2)
                    for x in shapes if x["geom"] == "roundRect" and x["box"][2] > EMU_IN * 0.3})
    borders = sorted({x["lw"] for x in shapes if x["geom"] in ("roundRect", "rect", "ellipse") and x["lw"]})
    arrows = [x for x in shapes if x["heads"]]
    aw = sorted({x["lw"] for x in arrows})
    heads = sorted({h for x in arrows for h in x["heads"]})
    dashes = sorted({x["dash"] for x in shapes if x["dash"]})
    issues = []
    if sharp:
        issues.append(f"{len(sharp)} 个非圆角模块（{', '.join(sorted({x['geom'] for x in sharp}))}）")
    if len(radii) > 1:
        issues.append(f"圆角半径 {len(radii)} 种 {radii} in")
    if len(borders) > 2:
        issues.append(f"边框粗细 {len(borders)} 种 {borders} pt")
    if len(aw) > 1:
        issues.append(f"箭头粗细 {len(aw)} 种 {aw} pt")
    if len(heads) > 1:
        issues.append(f"箭头样式 {heads}")
    if len(dashes) > 1:
        issues.append(f"虚线样式 {dashes}")
    rows.append(("R4 形状/线条统一", "⚠️" if issues else "✅",
                 "；".join(issues) if issues else f"圆角 {radii or '—'} in；边框 {borders or '—'} pt；箭头 {aw or '—'} pt {heads or ''}"))

    # R5 hues
    colors = [x["fill"] for x in shapes] + [x["line"] for x in shapes] + [r["color"] for r in runs]
    fam = Counter(hue_family(c) for c in colors if c and len(c) == 6)
    fam.pop(None, None)
    rows.append(("R5 色彩克制（2–4 种）", "❌" if len(fam) > 4 else "✅",
                 f"{len(fam)} 个色相：" + "，".join(f"{k}×{v}" for k, v in fam.most_common()) if fam else "只有灰阶"))

    # R6 collisions / out of frame
    labels = [x for x in shapes if x["text"]]
    clash = []
    for i in range(len(labels)):
        for j in range(i + 1, len(labels)):
            a, b = labels[i]["box"], labels[j]["box"]
            ix = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            iy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ix > 0.02 * EMU_IN and iy > 0.02 * EMU_IN:
                inside = lambda p, q: p[0] >= q[0] and p[1] >= q[1] and p[0] + p[2] <= q[0] + q[2] and p[1] + p[3] <= q[1] + q[3]
                if not (inside(a, b) or inside(b, a)):
                    clash.append(f"「{labels[i]['text'][:18]}」×「{labels[j]['text'][:18]}」")
    out = []
    if frame:
        fx, fy, fw, fh = frame["box"]
        out = [x["name"] for x in shapes if x["box"][2] and (x["box"][0] < fx - 1 or x["box"][1] < fy - 1 or
               x["box"][0] + x["box"][2] > fx + fw + 1 or x["box"][1] + x["box"][3] > fy + fh + 1)]
    crossed, tight = lines_through_text(slide), tight_gaps(slide)
    detail = ("文本框部分重叠：" + "；".join(clash[:4]) if clash else "无文本重叠") + (f"；{len(out)} 个元素超出 FRAME" if out else "")
    if crossed:
        detail += "；连线穿过文字：" + "；".join(f"「{t}」" for t in crossed[:4])
    if tight:
        detail += f"；{len(tight)} 处模块间距 < {T.GAP:.2f} in：" + "；".join(tight[:4]) + (" …" if len(tight) > 4 else "")
    rows.append(("R6 留白/不重叠", "⚠️" if clash or out or crossed or tight else "✅", detail))

    return rows, facts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--slides")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    theme = theme_fonts(prs)
    pick = {int(s) for s in a.slides.split(",")} if a.slides else None
    print(f"# 风格审计：{a.deck.split('/')[-1]}\n")
    fonts = {k: v for k, v in theme.items() if k.startswith("+")}
    print(f"幻灯片 {prs.slide_width / EMU_IN:.2f} × {prs.slide_height / EMU_IN:.2f} in；主题字体 {fonts or '—'}\n")
    for i, slide in enumerate(prs.slides, 1):
        if pick and i not in pick:
            continue
        rows, facts = audit_slide(slide, i, theme)
        print(f"## 第 {i} 页 · {facts['frame']} · 图片 {facts['pictures']} 张\n")
        print("| 规则 | 结果 | 细节 |\n|---|---|---|")
        for r in rows:
            print(f"| {r[0]} | {r[1]} | {r[2]} |")
        print()


if __name__ == "__main__":
    main()
